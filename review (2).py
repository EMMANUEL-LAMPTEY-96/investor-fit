"""Review a financial model.

Two independent checks:

1. scan_workbook(path): mechanical checks on any Excel model (needs openpyxl). Finds formula
   errors, hardcoded numbers inside forecast rows, formulas that break the pattern of their row,
   numbers buried inside formulas, links to other files, hidden sheets, and a balance sheet
   that does not balance.

2. assess_projections(data): judgement checks on the numbers themselves (plain Python). Is the
   growth believable, can the business repay the loan, does the amount asked match what the
   cash flow needs, does the use of funds add up?

Each finding has a severity (critical, high, medium, low), what was found, why a funder cares,
and how to fix it.
"""

from __future__ import annotations

import re
from collections import Counter
from typing import Any

ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#N/A", "#NAME?", "#NUM!", "#NULL!", "Err:")
SEVERITY_ORDER = {"critical": 0, "high": 1, "medium": 2, "low": 3}
# Numbers that are normal inside formulas: small counters/offsets, periods per year, unit scales.
HARMLESS_CONSTANTS = {str(i) for i in range(0, 11)} | {"12", "52", "100", "360", "365", "1000", "1000000", "0.5"}


def _finding(severity, area, where, found, why, fix):
    return {"severity": severity, "area": area, "where": where, "found": found,
            "why_it_matters": why, "fix": fix}


# ====================================================================== 1. mechanical scan

def scan_workbook(path: str, max_items_per_check: int = 15) -> dict[str, Any]:
    from openpyxl import load_workbook
    from openpyxl.formula.translate import Translator

    wb_f = load_workbook(path, data_only=False)
    wb_v = load_workbook(path, data_only=True)
    findings: list[dict] = []
    stats: dict[str, Any] = {"sheets": [], "formulas": 0, "numbers": 0}
    has_cached_values = False

    def add(*args):
        findings.append(_finding(*args))

    if getattr(wb_f, "_external_links", None):
        add("high", "Links", "Workbook", f"{len(wb_f._external_links)} link(s) to other files",
            "Numbers that depend on a file the funder does not have will break or go stale.",
            "Paste the linked values into an Inputs sheet and point formulas there.")

    for ws in wb_f.worksheets:
        wv = wb_v[ws.title]
        info = {"name": ws.title, "state": ws.sheet_state, "formulas": 0, "numbers": 0}
        if ws.sheet_state != "visible":
            add("low", "Transparency", ws.title, f"Sheet '{ws.title}' is hidden",
                "Funders and their analysts open hidden sheets; hidden logic looks like something is being concealed.",
                "Unhide it, or delete it if it is no longer used.")
        hidden_rows = sum(1 for d in ws.row_dimensions.values() if d.hidden)
        if hidden_rows:
            add("low", "Transparency", ws.title, f"{hidden_rows} hidden row(s)",
                "Hidden rows can hide hardcoded adjustments.", "Unhide them or group them with outline levels instead.")

        errors, magic, row_hardcodes, inconsistent = [], [], [], []
        for row in ws.iter_rows():
            formula_cells, number_cells = [], []
            for cell in row:
                v = cell.value
                if isinstance(v, str) and v.startswith("="):
                    info["formulas"] += 1
                    formula_cells.append(cell)
                    cached = wv[cell.coordinate].value
                    if cached is not None:
                        has_cached_values = True
                    if isinstance(cached, str) and cached.startswith(ERRORS):
                        errors.append(f"{cell.coordinate} ({cached})")
                    if re.search(r"\[[^\]]+\]", v):
                        errors.append(f"{cell.coordinate} (links to another file)")
                    body = re.sub(r'"[^"]*"', "", v)
                    body = re.sub(r"\$?[A-Z]{1,3}\$?\d+", "", body)
                    body = re.sub(r"'[^']+'!|[A-Za-z_]+!", "", body)
                    consts = [c for c in re.findall(r"(?<![A-Za-z_\d.])(\d+\.?\d*)", body) if c not in HARMLESS_CONSTANTS]
                    if consts:
                        magic.append(f"{cell.coordinate}: {v[:60]}")
                elif isinstance(v, (int, float)) and not isinstance(v, bool):
                    info["numbers"] += 1
                    number_cells.append(cell)

            # A typed-in number in the middle of a row of formulas = a likely overwritten forecast cell.
            if len(formula_cells) >= 2 and number_cells:
                first_f = min(c.column for c in formula_cells)
                for nc in number_cells:
                    if nc.column > first_f:
                        label = ws.cell(row=nc.row, column=1).value or ws.cell(row=nc.row, column=2).value
                        row_hardcodes.append(f"{nc.coordinate}" + (f" ('{label}')" if label else ""))

            # Formulas that break the pattern of their row. The first formula in a row is skipped:
            # first periods and opening balances legitimately differ from later periods.
            if len(formula_cells) >= 4:
                canon = []
                cells = formula_cells[1:]
                anchor = cells[0]
                for c in cells:
                    try:
                        canon.append(Translator(c.value, origin=c.coordinate).translate_formula(anchor.coordinate))
                    except Exception:
                        canon.append(c.value)
                common, count = Counter(canon).most_common(1)[0]
                if count >= 2 and count < len(canon):
                    for c, k in zip(cells, canon):
                        if k != common:
                            label = ws.cell(row=c.row, column=1).value
                            inconsistent.append(f"{c.coordinate}" + (f" ('{label}')" if label else ""))

        def trim(xs):
            return ", ".join(xs[:max_items_per_check]) + (f" and {len(xs) - max_items_per_check} more" if len(xs) > max_items_per_check else "")

        if errors:
            add("critical", "Errors", ws.title, f"{len(errors)} cell(s) with errors: {trim(errors)}",
                "A funder who sees #REF! or #DIV/0! stops trusting every other number in the model.",
                "Trace each error to its source; guard divisions with IF(denominator=0, ...).")
        if row_hardcodes:
            add("high", "Hardcoded numbers", ws.title,
                f"{len(row_hardcodes)} typed-in number(s) inside rows of formulas: {trim(row_hardcodes)}",
                "A forecast cell typed over a formula no longer responds to the assumptions, so the model "
                "silently gives the wrong answer when anything changes.",
                "Replace each with the row's formula, or move the number to the assumptions and reference it.")
        if inconsistent:
            add("high", "Inconsistent formulas", ws.title,
                f"{len(inconsistent)} formula(s) that differ from the rest of their row: {trim(inconsistent)}",
                "One edited cell in a projection row is the most common silent error in financial models.",
                "Check each flagged cell; if it should match its neighbours, copy the formula across.")
        if magic:
            add("medium", "Numbers inside formulas", ws.title,
                f"{len(magic)} formula(s) with typed-in numbers: {trim(magic)}",
                "Assumptions hidden inside formulas (like *1.15) can't be seen, questioned or changed by a funder.",
                "Put each number in a labelled assumption cell and reference it.")
        stats["formulas"] += info["formulas"]
        stats["numbers"] += info["numbers"]
        stats["sheets"].append(info)

    # Balance sheet check: find 'total assets' and 'total liabilities ... equity' rows and compare.
    bs = _balance_check(wb_v)
    if bs:
        findings.append(bs)

    if stats["formulas"] == 0:
        add("critical", "Structure", "Workbook", "No formulas at all; every number is typed in",
            "A model with no formulas is a table of hopes. A funder cannot test what happens if sales are 20% lower.",
            "Rebuild it with assumptions driving formulas (the agent can build one for you).")
    if stats["formulas"] and not has_cached_values:
        add("low", "Structure", "Workbook", "Formula results are not saved in the file",
            "Some checks could not read the calculated numbers.",
            "Open the file in Excel, save it, and run the review again.")

    findings.sort(key=lambda f: SEVERITY_ORDER[f["severity"]])
    return {"stats": stats, "findings": findings}


def _balance_check(wb_v) -> dict | None:
    for ws in wb_v.worksheets:
        ta_row = tle_row = None
        label_col = 1
        for row in ws.iter_rows(max_col=3):
            for cell in row:
                if not isinstance(cell.value, str):
                    continue
                t = cell.value.strip().lower()
                if t.startswith("total assets"):
                    ta_row, label_col = cell.row, cell.column
                elif "total" in t and "liabilit" in t and "equity" in t:
                    tle_row = cell.row
        if ta_row and tle_row:
            diffs = []
            for c in range(label_col + 1, ws.max_column + 1):
                a, b = ws.cell(ta_row, c).value, ws.cell(tle_row, c).value
                if isinstance(a, (int, float)) and isinstance(b, (int, float)) and abs(a - b) > max(1, abs(a) * 1e-6):
                    diffs.append(f"{ws.cell(ta_row, c).coordinate}: difference {a - b:,.0f}")
            if diffs:
                return _finding("critical", "Balance sheet", ws.title,
                                f"Balance sheet does not balance in {len(diffs)} period(s): {', '.join(diffs[:6])}",
                                "An unbalanced balance sheet means the statements are not linked correctly; "
                                "the cash figure, and so the funding need, cannot be trusted.",
                                "Check that closing cash comes from the cash flow statement and that equity "
                                "includes each year's profit and any new equity.")
            return None
    return None


# ====================================================================== 2. judgement checks

def _cagr(first, last, years):
    if not first or first <= 0 or not last or last <= 0 or years <= 0:
        return None
    return (last / first) ** (1 / years) - 1


def assess_projections(d: dict[str, Any]) -> dict[str, Any]:
    """d: years, revenue, ebitda, net_income, closing_cash, total_debt, optional debt_service,
    optional history {revenue, ebitda}, optional ask {amount, instrument, tenor_years, interest_rate,
    grace_years, use_of_funds_total}, optional min_dscr (default 1.25)."""
    f: list[dict] = []
    years = d.get("years") or []
    n = len(d.get("revenue") or [])
    rev = d.get("revenue") or []
    ebitda = d.get("ebitda") or []
    ni = d.get("net_income") or []
    cash = d.get("closing_cash") or []
    debt = d.get("total_debt") or []
    hist = d.get("history") or {}
    ask = d.get("ask") or {}
    min_dscr = float(d.get("min_dscr") or 1.25)
    cur = d.get("currency") or ""
    yl = lambda i: str(years[i]) if i < len(years) else f"year {i + 1}"  # noqa: E731
    money = lambda x: f"{cur} {x:,.0f}".strip()  # noqa: E731
    metrics: dict[str, Any] = {}

    if n and n < 3:
        f.append(_finding("medium", "Horizon", "Projections", f"Only {n} year(s) projected",
                          "Lenders and investors usually expect 3 to 5 years, and at least the full loan tenor.",
                          "Extend the projections to cover the whole loan period."))

    # Growth
    if n >= 2:
        g = _cagr(rev[0], rev[-1], n - 1)
        metrics["revenue_cagr"] = g
        if g is not None and g > 0.5:
            f.append(_finding("high", "Growth", "Revenue", f"Revenue grows {g:.0%} a year on average ({yl(0)}–{yl(n - 1)})",
                              "Growth this fast is rare for an established SME; funders will discount the whole plan.",
                              "Tie growth to evidence: signed contracts, capacity, pipeline, market share. Show a slower case."))
        elif g is not None and g > 0.3:
            f.append(_finding("medium", "Growth", "Revenue", f"Revenue grows {g:.0%} a year on average",
                              "Ambitious; funders will ask what drives it.",
                              "Show the drivers (volume, price, new customers) and a downside case."))
        for i in range(1, n):
            if rev[i - 1] and rev[i] / rev[i - 1] > 2:
                f.append(_finding("high", "Growth", yl(i), f"Revenue more than doubles in {yl(i)}",
                                  "Single-year jumps are the first thing an analyst questions.",
                                  "Explain the step change (new plant, new contract) or smooth it."))
    hrev = hist.get("revenue") or []
    if len(hrev) >= 2 and n >= 1:
        hg = _cagr(hrev[0], hrev[-1], len(hrev) - 1)
        pg = (rev[0] / hrev[-1] - 1) if hrev[-1] else None
        metrics["historical_growth"] = hg
        if hg is not None and pg is not None and pg > max(0.15, 2 * hg):
            f.append(_finding("high", "Growth", "First forecast year",
                              f"First-year growth of {pg:.0%} versus {hg:.0%} a year historically",
                              "A forecast that breaks sharply from the track record needs a clear reason.",
                              "Explain what changes (the funding itself, a new contract) or bring it closer to history."))

    # Margins
    margins = [(e / r) if r else None for e, r in zip(ebitda, rev)]
    metrics["ebitda_margin"] = margins
    valid = [m for m in margins if m is not None]
    if len(valid) >= 2 and valid[-1] - valid[0] > 0.10:
        f.append(_finding("medium", "Margins", "EBITDA margin",
                          f"EBITDA margin rises from {valid[0]:.0%} to {valid[-1]:.0%}",
                          "Rising margins assume costs grow slower than sales. Funders want to know why.",
                          "Justify with scale effects, pricing or efficiency gains, or hold margins flatter."))
    hm = hist.get("ebitda_margin")
    if hm is None and hist.get("ebitda") and hrev:
        hm = hist["ebitda"][-1] / hrev[-1] if hrev[-1] else None
    if hm is not None and valid and valid[0] - hm > 0.08:
        f.append(_finding("high", "Margins", "First forecast year",
                          f"Forecast EBITDA margin {valid[0]:.0%} versus {hm:.0%} last year",
                          "A margin jump right after funding is a red flag unless the funding causes it.",
                          "Show the cost savings or price changes behind it."))
    if any(m is not None and m > 0.45 for m in margins):
        f.append(_finding("low", "Margins", "EBITDA margin", "EBITDA margin above 45% in at least one year",
                          "Unusually high for most SMEs outside software.", "Compare with sector benchmarks."))

    # Profitability and cash
    if ni and ni[-1] is not None and ni[-1] < 0:
        f.append(_finding("high", "Profitability", yl(n - 1), "Still loss-making in the final projection year",
                          "Lenders need profits to be repaid; equity investors need a visible path to profit.",
                          "Show when and how the business reaches break-even."))
    if cash and min(cash) < 0:
        i = cash.index(min(cash))
        f.append(_finding("critical", "Cash", yl(i), f"Cash goes negative ({money(min(cash))}) in {yl(i)} even with the funding",
                          "The plan as written runs out of money; the amount raised is too small or too late.",
                          "Raise more, phase spending, or improve working capital (collect faster, pay later)."))

    # Debt capacity
    ds = d.get("debt_service")
    amount = float(ask.get("amount") or 0)
    instrument = (ask.get("instrument") or "").lower()
    if not ds and amount and instrument in ("debt", "loan", "working_capital", "trade_finance"):
        tenor = int(ask.get("tenor_years") or 5)
        grace = int(ask.get("grace_years") or 0)
        rate = float(ask.get("interest_rate") or 0.12)
        bal, ds = amount, []
        for t in range(1, n + 1):
            principal = amount / max(1, tenor - grace) if grace < t <= tenor else 0
            principal = min(bal, principal)
            ds.append(bal * rate + principal)
            bal -= principal
        metrics["debt_service_estimated"] = True
    if ds and ebitda:
        dscr = [(e / s) if s else None for e, s in zip(ebitda, ds)]
        metrics["dscr"] = dscr
        low = [(i, x) for i, x in enumerate(dscr) if x is not None and x < min_dscr]
        if low:
            worst = min(low, key=lambda p: p[1])
            sev = "critical" if worst[1] < 1 else "high"
            f.append(_finding(sev, "Debt capacity", yl(worst[0]),
                              f"Debt service cover falls to {worst[1]:.2f}x in {yl(worst[0])} (minimum {min_dscr:.2f}x)"
                              + (" — estimated from the loan terms" if metrics.get("debt_service_estimated") else ""),
                              "Lenders test whether cash profit covers interest and repayments with a buffer. "
                              "Below 1.0x the loan cannot be repaid from operations.",
                              "Lengthen the tenor, add a grace period, borrow less, or mix in equity."))
    if debt and ebitda:
        lev = [(dbt / e) if e and e > 0 else None for dbt, e in zip(debt, ebitda)]
        metrics["debt_to_ebitda"] = lev
        if lev and lev[0] is not None and lev[0] > 4:
            f.append(_finding("high", "Leverage", yl(0), f"Debt is {lev[0]:.1f}x EBITDA after funding",
                              "Most SME lenders are uncomfortable above 3–4x.",
                              "Reduce the loan or add equity."))

    # Ask versus need
    if amount and cash:
        low_cash = min(cash)
        metrics["lowest_cash_with_funding"] = low_cash
        if low_cash > 0.6 * amount:
            f.append(_finding("medium", "Size of the ask", "Funding",
                              f"Cash never falls below {money(low_cash)}, which is {low_cash / amount:.0%} of the {money(amount)} being raised",
                              "Raising more than the plan needs costs interest or dilution, and funders notice idle cash.",
                              "Size the raise to the peak funding need plus a buffer, or explain the reserve."))
    uof = ask.get("use_of_funds_total")
    if amount and uof is not None and abs(uof - amount) > 0.02 * amount:
        f.append(_finding("high", "Use of funds", "Funding",
                          f"Use of funds adds up to {money(uof)} but the ask is {money(amount)}",
                          "Every amount requested must be accounted for; a mismatch suggests the plan is not finished.",
                          "Make the use-of-funds table add up exactly to the amount requested."))
    if amount and uof is None:
        f.append(_finding("medium", "Use of funds", "Funding", "No use-of-funds breakdown given",
                          "Funders always ask exactly what the money will pay for.",
                          "Add a table: item, amount, timing."))

    f.sort(key=lambda x: SEVERITY_ORDER[x["severity"]])
    counts = Counter(x["severity"] for x in f)
    verdict = ("Not ready to send: fix the critical issues first." if counts.get("critical")
               else "Needs work before sending to funders." if counts.get("high")
               else "Broadly fundable; tidy up the points below." if f
               else "No issues found by these checks.")
    return {"verdict": verdict, "counts": dict(counts), "findings": f, "metrics": metrics}


def from_engine(result: dict[str, Any]) -> dict[str, Any]:
    """Turn engine.run() output into the input of assess_projections()."""
    a = result["assumptions"]
    is_, bs = result["income_statement"], result["balance_sheet"]
    ds = [result["debt"]["new"]["interest"][i] + result["debt"]["existing"]["interest"][i]
          + result["debt"]["new"]["repayment"][i] + result["debt"]["existing"]["repayment"][i]
          for i in range(a["years"])]
    ebitda_after_tax = [is_["ebitda"][i] - is_["tax"][i] for i in range(a["years"])]
    instrument = "debt" if a["new_debt"]["amount"] >= a["new_equity"]["amount"] else "equity"
    return {
        "currency": a["currency"], "years": result["years"], "revenue": is_["revenue"],
        "ebitda": ebitda_after_tax, "net_income": is_["net_income"], "closing_cash": bs["cash"],
        "total_debt": bs["debt"], "debt_service": ds if any(ds) else None,
        "ask": {"amount": a["new_debt"]["amount"] + a["new_equity"]["amount"], "instrument": instrument,
                "use_of_funds_total": sum(u["amount"] for u in a["use_of_funds"]) if a["use_of_funds"] else None},
        "min_dscr": a["minimum_dscr"],
    }
