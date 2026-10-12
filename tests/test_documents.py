import copy
import json
import subprocess
import sys
from pathlib import Path

import pytest

from investor_fit.evidence import bank_summary, check
from investor_fit.finmodel.engine import run
from investor_fit.finmodel.from_evidence import derive

ROOT = Path(__file__).resolve().parent.parent
EXAMPLE = json.loads((ROOT / "skill" / "investor-fit-advisor" / "examples" / "evidence_example.json").read_text())
ANSWERS = {"existing_debt": {"remaining_years": 3},
           "new_debt": {"amount": 500000, "currency": "USD", "interest_rate": 0.11, "tenor_years": 5,
                        "grace_years": 1, "drawdown_year": 1},
           "fx": {"USD": {"rate": 12.0, "depreciation": 0.08}},
           "use_of_funds": [{"item": "Processing line", "amount": 3000000}, {"item": "Stock", "amount": 3000000}],
           "capex": [{"name": "Processing line", "amounts": [3000000, 0, 0, 0, 0]}]}


def issues(report):
    return [x["issue"] for x in report["items"]]


def test_finds_the_planted_problems():
    r = check(EXAMPLE)
    found = " | ".join(issues(r))
    for expected in ("Balance sheet does not balance in 2025", "Revenue 2025 differs between documents",
                     "has expired", "is not signed", "Equity movement in 2025", "No deposits in 2025-06",
                     "overdrawn", "Plan revenue for 2027"):
        assert expected in found, expected
    assert r["counts_by_severity"]["critical"] == 1
    assert r["questions"] and all(":" in q for q in r["questions"])


def test_clean_documents_raise_no_inconsistencies():
    ev = copy.deepcopy(EXAMPLE)
    ev["facts"] = [f for f in ev["facts"] if f["document"] != "plan"]
    for f in ev["facts"]:
        if f["metric"] == "equity" and f["period"] == "2025":
            f["value"] = 6480  # assets 8,640 = payables 960 + debt 1,200 + equity 6,480
    # equity 6,100 + profit 682 - dividend 302 = 6,480
    ev["facts"].append({"metric": "dividends", "period": "2025", "value": 302, "document": "afs2025"})
    r = check(ev)
    assert not [x for x in r["items"] if x["category"] == "inconsistent"]


def test_unit_mix_up_is_spotted():
    ev = copy.deepcopy(EXAMPLE)
    for f in ev["facts"]:
        if f["metric"] == "inventory" and f["period"] == "2025":
            f["value"] = 1800000  # entered in units although the document is in thousands
    assert any("different unit" in i for i in issues(check(ev)))


def test_bank_transactions_are_summarised_and_flagged():
    txs = [{"date": f"2025-{m:02d}-05", "description": "Sales receipt", "amount": 100000} for m in range(1, 13)]
    txs += [{"date": "2025-03-09", "description": "Cheque returned unpaid", "amount": -40000}]
    b = bank_summary({"transactions": txs})
    assert b["mode"] == "transactions" and b["months_covered"] == 12
    assert any("returned or unpaid" in f["issue"] for f in b["flags"])


def test_low_bank_deposits_against_revenue_are_questioned():
    ev = copy.deepcopy(EXAMPLE)
    for m in ev["bank"]["months"]:
        m["deposits"] = m["deposits"] / 3
    assert any("Bank deposits are only" in i for i in issues(check(ev)))


def test_asks_before_building_and_builds_after_answers():
    first = derive(EXAMPLE)
    assert first["status"] == "needs_input" and first["assumptions"] is None
    fields = {q["field"] for q in first["questions"]}
    assert {"new_debt", "use_of_funds", "capex", "existing_debt.remaining_years"} <= fields
    done = derive(EXAMPLE, ANSWERS)
    assert done["status"] == "ready"
    a = done["assumptions"]
    origins = {r["field"]: r["origin"] for r in done["register"]}
    assert origins["opening_balance.cash"] == "document"
    assert origins["cogs_pct_revenue"] == "derived" and origins["new_debt.amount"] == "user"
    assert a["opening_balance"]["cash"] == 600000  # scaled from thousands
    assert a["cogs_pct_revenue"] == pytest.approx((7800 / 12100 + 9100 / 14200) / 2, abs=1e-4)
    r = run(a)
    assert all(abs(x) < 0.01 for x in r["balance_sheet"]["balance_check"])


def test_placeholders_only_when_allowed_and_are_labelled():
    r = derive(EXAMPLE, allow_placeholders=True)
    assert r["status"] == "ready" and r["placeholders_used"]
    assert all(x["origin"] == "placeholder" for x in r["register"] if x["field"] in r["placeholders_used"])


def test_scripts_run_from_packaged_skill(tmp_path):
    import zipfile

    subprocess.run([sys.executable, str(ROOT / "scripts" / "package_skill.py")], check=True, capture_output=True)
    with zipfile.ZipFile(ROOT / "dist" / "investor-fit-advisor.zip") as z:
        z.extractall(tmp_path)
    s = tmp_path / "investor-fit-advisor" / "scripts"
    ev = tmp_path / "investor-fit-advisor" / "examples" / "evidence_example.json"
    (tmp_path / "answers.json").write_text(json.dumps(ANSWERS))
    out = json.loads(subprocess.run([sys.executable, "-I", str(s / "build_from_documents.py"), str(ev), "--out",
                                     str(tmp_path / "m.xlsx")], capture_output=True, text=True, check=True).stdout)
    assert out["status"] == "needs_input"
    out = json.loads(subprocess.run([sys.executable, "-I", str(s / "build_from_documents.py"), str(ev), "--out",
                                     str(tmp_path / "m.xlsx"), "--answers", str(tmp_path / "answers.json")],
                                    capture_output=True, text=True, check=True).stdout)
    assert out["status"] == "built"
    import openpyxl

    names = openpyxl.load_workbook(tmp_path / "m.xlsx").sheetnames
    assert {"Data quality", "Sources", "Scenarios", "Stress tests"} <= set(names)
